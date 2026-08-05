"""Tests for /api/export — exercises the real templates via excel_service."""

import io

from openpyxl import load_workbook

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_export_cmr_fills_template(client):
    payload = {
        "doc_type": "cmr",
        "fields": {
            "shipper": "SC Exemplu SRL",
            "cargo_description": "Piese auto",
            "gross_weight_kg": 450,
        },
    }
    resp = client.post("/api/export", json=payload)
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(XLSX)
    assert "cmr_export.xlsx" in resp.headers.get("content-disposition", "")
    assert resp.content[:2] == b"PK"  # xlsx is a zip

    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.active.title == "CMR"
    sheet, coord = list(wb.defined_names["shipper"].destinations)[0]
    assert wb[sheet][coord].value == "SC Exemplu SRL"


def test_export_awb(client):
    resp = client.post(
        "/api/export",
        json={"doc_type": "awb", "fields": {"shipper_name": "X SRL"}},
    )
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.active.title == "AWB"
    sheet, coord = list(wb.defined_names["shipper_name"].destinations)[0]
    assert wb[sheet][coord].value == "X SRL"


def test_export_rejects_removed_bol(client):
    resp = client.post("/api/export", json={"doc_type": "bol", "fields": {}})
    assert resp.status_code == 422


def test_export_requires_doc_type(client):
    resp = client.post("/api/export", json={"fields": {}})
    assert resp.status_code == 422


def test_export_ignores_non_template_fields(client):
    # `confidence` isn't a named range — it must be ignored, not error.
    resp = client.post(
        "/api/export",
        json={"doc_type": "cmr", "fields": {"confidence": "high", "shipper": "A"}},
    )
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    assert "confidence" not in wb.defined_names


def test_export_neutralizes_formula_injection(client):
    # A value from a third-party document that would be a live formula must be
    # written as literal text (leading "'"), not evaluated on open.
    payload = {
        "doc_type": "cmr",
        "fields": {"cargo_description": '=HYPERLINK("http://evil","x")'},
    }
    resp = client.post("/api/export", json=payload)
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    sheet, coord = list(wb.defined_names["cargo_description"].destinations)[0]
    cell = wb[sheet][coord]
    assert cell.value.startswith("'=")  # quoted → inert text, not a formula
    assert cell.data_type == "s"  # stored as string, not a formula
