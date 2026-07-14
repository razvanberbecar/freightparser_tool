"""Generate the CMR and AWB Excel templates with named ranges (Phase 4).

Run from the backend directory:
    ./venv/Scripts/python.exe scripts/build_templates.py

Produces backend/templates/cmr_template.xlsx and awb_template.xlsx. Each
fillable cell has a workbook-level named range whose name matches the JSON keys
Claude returns, so excel_service can fill by name.

- CMR layout recreates the trilingual (RO/EN/FR) Romanian CMR from the source
  .xls (255655453-CMR-MODEL.xls).
- AWB layout recreates the IATA Air Waybill from the source PDF.
"""

import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

TEMPLATES_DIR = os.path.join(os.path.dirname(__file__), "..", "templates")

LABEL_FONT = Font(size=7)
TITLE_FONT = Font(size=12, bold=True)
FILL_FONT = Font(size=9)
FILL_STYLE = PatternFill("solid", fgColor="FFFFF7CC")  # pale yellow = fillable
_THIN = Side(style="thin")


def _merge_border(existing: Border, **sides) -> Border:
    """Return a Border that adds the given sides without dropping existing ones."""
    return Border(
        top=sides.get("top", existing.top),
        bottom=sides.get("bottom", existing.bottom),
        left=sides.get("left", existing.left),
        right=sides.get("right", existing.right),
    )


def outline(ws, r1, c1, r2, c2):
    """Draw a thin border around the rectangle (0-based, inclusive)."""
    for c in range(c1, c2 + 1):
        top = ws.cell(row=r1 + 1, column=c + 1)
        top.border = _merge_border(top.border, top=_THIN)
        bot = ws.cell(row=r2 + 1, column=c + 1)
        bot.border = _merge_border(bot.border, bottom=_THIN)
    for r in range(r1, r2 + 1):
        left = ws.cell(row=r + 1, column=c1 + 1)
        left.border = _merge_border(left.border, left=_THIN)
        right = ws.cell(row=r + 1, column=c2 + 1)
        right.border = _merge_border(right.border, right=_THIN)


def label(ws, r, c, text, font=LABEL_FONT):
    cell = ws.cell(row=r + 1, column=c + 1, value=text)
    cell.font = font
    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)


def add_fill(ws, wb, sheet, name, r1, c1, r2, c2):
    """Merge a fill region, style it, and register a named range at its anchor."""
    if (r1, c1) != (r2, c2):
        ws.merge_cells(
            start_row=r1 + 1, start_column=c1 + 1, end_row=r2 + 1, end_column=c2 + 1
        )
    anchor = ws.cell(row=r1 + 1, column=c1 + 1)
    anchor.font = FILL_FONT
    anchor.fill = FILL_STYLE
    anchor.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    ref = f"{sheet}!${get_column_letter(c1 + 1)}${r1 + 1}"
    wb.defined_names.add(DefinedName(name, attr_text=ref))


# --------------------------------------------------------------------------- #
# CMR
# --------------------------------------------------------------------------- #

# (row, col, text) — verbatim from the source .xls dump.
CMR_LABELS = [
    (0, 0, "1. Expeditor (denumire, adresa, tara)"),
    (1, 0, "Sender (name, address, country)"),
    (2, 0, "Expediteur (nom, adresse, pays)"),
    (0, 9, "SCRISOARE DE TRANSPORT"),
    (1, 9, "CONSIGNMENT NOTE"),
    (2, 9, "LETTRE DE VOITURE"),
    (3, 13, "(CMR)"),
    (6, 0, "2. Destinatar (nume, adresa, tara)"),
    (7, 0, "Consignee (name, address, country)"),
    (8, 0, "Destinataire (nom, adresse, pays)"),
    (6, 9, "16. Operator de transport (denumire, adresa, tara)"),
    (7, 9, "Carrier (name, address, country)"),
    (8, 9, "Transporteur (nom, adresse, pays)"),
    (12, 0, "3. Locul descarcarii (loc, tara)"),
    (13, 0, "Place of delivery of goods (place, country)"),
    (14, 0, "Lieu prevu pour la livraison (lieu, pays)"),
    (12, 9, "17. Transportatori succesivi (nume, adresa, tara)"),
    (13, 9, "Successive carriers (name, address, country)"),
    (14, 9, "Transporteurs successifs (nom, adresse, pays)"),
    (17, 0, "4. Locul incarcarii (loc, tara, data)"),
    (18, 0, "Place and date of taking over the goods"),
    (19, 0, "Lieu et date de la prise en charge"),
    (17, 9, "18. Rezerve si observatii ale transportatorilor"),
    (18, 9, "Carrier's reservation and observations"),
    (19, 9, "Reserves et observations du transporteur"),
    (22, 0, "5. Documente anexate"),
    (23, 0, "Documents attached"),
    (24, 0, "Documents annexes"),
    (26, 0, "6. Marci si numere"),
    (27, 0, "Marks and Nos"),
    (28, 0, "Marques et numeros"),
    (26, 3, "7. Nr. de colete"),
    (27, 3, "Number of packages"),
    (28, 3, "Nombre de colis"),
    (26, 6, "8. Mod de ambalare"),
    (27, 6, "Method of packing"),
    (28, 6, "Mode d'emballage"),
    (26, 8, "9. Natura marfii"),
    (27, 8, "Nature of the goods"),
    (28, 8, "Nature de la marchandise"),
    (26, 13, "10. Numar statistic"),
    (27, 13, "Statistical number"),
    (28, 13, "No statistique"),
    (26, 15, "11. Greutate bruta kg"),
    (27, 15, "Gross weight kg"),
    (28, 15, "Poids brut kg"),
    (26, 17, "12. Cubaj m3"),
    (27, 17, "Volume m3"),
    (28, 17, "Cubage m3"),
    (37, 0, "Clasa / Classe"),
    (37, 3, "Cifra / Number"),
    (37, 6, "Litera / Letter"),
    (37, 8, "(ADR*)"),
    (40, 0, "13. Instructiunile expeditorului"),
    (41, 0, "Sender's instructions"),
    (42, 0, "Instructions de l'expediteur"),
    (40, 9, "19. Conventii speciale"),
    (41, 9, "Special agreements"),
    (42, 9, "Conventions particulieres"),
    (46, 9, "20. Plata prin / To be paid by / A payer par"),
    (56, 0, "14. Instructiuni de plata / Instructions as to payment"),
    (57, 0, "Prescriptions d'affranchissement"),
    (58, 2, "Franco / Carriage paid / Plata la expediere"),
    (59, 2, "Non franco / Carriage forward / Plata la destinatie"),
    (60, 0, "21. Stabilit in / Established in / Etabli a"),
    (60, 4, "data / on / le"),
    (60, 9, "15. Suma de plata / Cash on delivery"),
    (61, 9, "Remboursement"),
    (63, 0, "22. Semnatura si stampila expeditorului"),
    (63, 6, "23. Semnatura si stampila transportatorului"),
    (63, 13, "24. Receptia marfii / Goods received"),
]

# (name, r1, c1, r2, c2) fill regions.
CMR_FILLS = [
    ("shipper", 3, 0, 5, 8),
    ("consignee", 9, 0, 11, 8),
    ("delivery_place", 15, 0, 16, 8),
    ("loading_place", 20, 0, 20, 8),
    ("loading_date", 21, 0, 21, 8),
    ("attached_docs", 25, 0, 25, 8),
    ("marks", 29, 0, 36, 2),
    ("num_packages", 29, 3, 36, 5),
    ("packing_method", 29, 6, 36, 7),
    ("cargo_description", 29, 8, 36, 12),
    ("statistical_number", 29, 13, 36, 14),
    ("gross_weight_kg", 29, 15, 36, 16),
    ("volume_m3", 29, 17, 36, 17),
    ("adr_class", 38, 0, 39, 8),
    ("sender_instructions", 43, 0, 55, 8),
    ("payment_instructions", 58, 0, 59, 1),
    ("carrier", 9, 9, 11, 17),
    ("special_agreements", 43, 9, 45, 17),
    ("established_place", 60, 2, 60, 3),
    ("established_date", 60, 5, 60, 8),
    ("cash_on_delivery", 62, 9, 62, 17),
]

CMR_BOXES = [
    (0, 0, 5, 8), (6, 0, 11, 8), (12, 0, 16, 8), (17, 0, 21, 8), (22, 0, 25, 8),
    (37, 0, 39, 8), (40, 0, 55, 8), (56, 0, 59, 8), (60, 0, 62, 8),
    (63, 0, 69, 5), (63, 6, 69, 12), (63, 13, 69, 17),
    (0, 9, 5, 17), (6, 9, 11, 17), (12, 9, 16, 17), (17, 9, 21, 17),
    (40, 9, 45, 17), (46, 9, 59, 17), (60, 9, 62, 17),
    # goods table columns (row band 26-36 spans full width)
    (26, 0, 36, 2), (26, 3, 36, 5), (26, 6, 36, 7), (26, 8, 36, 12),
    (26, 13, 36, 14), (26, 15, 36, 16), (26, 17, 36, 17),
]


def build_cmr(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "CMR"
    for c in range(18):
        ws.column_dimensions[get_column_letter(c + 1)].width = 9.5
    for r, c, text in CMR_LABELS:
        label(ws, r, c, text, font=TITLE_FONT if (r <= 2 and c == 9) else LABEL_FONT)
    for box in CMR_BOXES:
        outline(ws, *box)
    for name, r1, c1, r2, c2 in CMR_FILLS:
        add_fill(ws, wb, "CMR", name, r1, c1, r2, c2)
    wb.save(path)


# --------------------------------------------------------------------------- #
# AWB
# --------------------------------------------------------------------------- #

AWB_LABELS = [
    (0, 0, "Shipper's Name and Address"),
    (0, 7, "Not Negotiable"),
    (1, 7, "Air Waybill"),
    (2, 7, "Issued By"),
    (3, 7, "Copies 1, 2 and 3 of this Air Waybill are originals and have the same validity"),
    (5, 0, "Consignee's Name and Address"),
    (5, 7, "It is agreed that the goods described herein are accepted in apparent good order "
           "and condition (except as noted) for carriage SUBJECT TO THE CONDITIONS OF "
           "CONTRACT ON THE REVERSE HEREOF."),
    (12, 0, "Issuing Carrier's Agent Name and City"),
    (12, 7, "Accounting Information:"),
    (17, 0, "Agent's IATA Code"),
    (17, 4, "Account No."),
    (19, 0, "Airport of Departure (Addr. of First Carrier) and Requested Routing"),
    (19, 7, "Reference Number"),
    (19, 10, "Optional Shipping Information"),
    (21, 0, "To"),
    (21, 1, "By First Carrier"),
    (21, 3, "to"),
    (21, 4, "by"),
    (21, 5, "to"),
    (21, 6, "by"),
    (21, 7, "Currency"),
    (21, 8, "CHGS Code"),
    (21, 9, "WT/VAL PPD COLL"),
    (21, 10, "Other PPD COLL"),
    (21, 11, "Declared Value for Carriage"),
    (21, 13, "Declared Value for Customs"),
    (23, 0, "Airport of Destination"),
    (23, 3, "Flight Date"),
    (23, 4, "For Carrier Use Only"),
    (23, 5, "Flight Date"),
    (23, 7, "Amount of Insurance"),
    (23, 9, "INSURANCE - If carrier offers insurance, and such insurance is requested, "
            "indicate amount to be insured in figures in box marked \"Amount of Insurance\"."),
    (26, 0, "Handling Information"),
    (27, 12, "SCI"),
    (28, 0, "No. of Pieces RCP"),
    (28, 1, "Gross Weight"),
    (28, 2, "kg/lb"),
    (28, 3, "Rate Class / Commodity Item No."),
    (28, 4, "Chargeable Weight"),
    (28, 5, "Rate / Charge"),
    (28, 6, "Total"),
    (28, 7, "Nature and Quantity of Goods (inc. Dimensions or Volume)"),
    (41, 0, "Prepaid"),
    (41, 1, "Weight Charge"),
    (41, 2, "Collect"),
    (42, 1, "Valuation Change"),
    (43, 1, "Tax"),
    (44, 1, "Total Other Charges Due Agent"),
    (45, 1, "Total Other Charges Due Carrier"),
    (46, 0, "Total Prepaid"),
    (46, 3, "Total Collect"),
    (47, 0, "Currency Conversion Rates"),
    (47, 3, "CC Charges in Dest. Currency"),
    (48, 0, "For Carrier's Use only at Destination"),
    (48, 3, "Charges at Destination"),
    (48, 5, "Total Collect Charges"),
    (41, 7, "Other Charges"),
    (44, 7, "Shipper certifies that the particulars on the face hereof are correct and that "
            "insofar as any part of the consignment contains dangerous goods, such part is "
            "properly described by name and is in proper condition for carriage by air "
            "according to the applicable Dangerous Goods Regulations."),
    (46, 7, "Signature of Shipper or his Agent"),
    (47, 7, "Executed on (date)"),
    (47, 10, "at (place)"),
    (47, 12, "Signature of Issuing Carrier or its Agent"),
]

AWB_FILLS = [
    ("shipper_name", 1, 0, 1, 6),
    ("shipper_address", 2, 0, 4, 6),
    ("consignee_name", 6, 0, 6, 6),
    ("consignee_address", 7, 0, 11, 6),
    ("agent_name", 13, 0, 13, 6),
    ("agent_city", 14, 0, 16, 6),
    ("accounting_info", 13, 7, 18, 13),
    ("airport_departure", 20, 0, 20, 6),
    ("reference_number", 20, 7, 20, 9),
    ("flight_number", 22, 1, 22, 2),
    ("declared_value_carriage", 22, 11, 22, 12),
    ("declared_value_customs", 22, 13, 22, 13),
    ("airport_destination", 24, 0, 24, 2),
    ("flight_date", 24, 3, 24, 3),
    ("insurance_amount", 24, 7, 24, 8),
    ("handling_info", 27, 0, 27, 11),
    ("num_pieces", 29, 0, 40, 0),
    ("gross_weight_kg", 29, 1, 40, 1),
    ("rate_class", 29, 3, 40, 3),
    ("chargeable_weight", 29, 4, 40, 4),
    ("cargo_description", 29, 7, 38, 13),
    ("cargo_dimensions", 39, 7, 40, 13),
    ("other_charges", 42, 7, 43, 13),
    ("total_prepaid", 46, 0, 46, 2),
    ("total_collect", 46, 3, 46, 6),
    ("execution_date", 47, 8, 47, 9),
    ("execution_place", 47, 11, 47, 11),
]

AWB_BOXES = [
    (0, 0, 4, 6), (0, 7, 2, 13), (3, 7, 4, 13),
    (5, 0, 11, 6), (5, 7, 11, 13),
    (12, 0, 16, 6), (12, 7, 18, 13), (17, 0, 17, 3), (17, 4, 18, 6),
    (19, 0, 22, 6), (19, 7, 20, 13), (21, 7, 22, 13),
    (23, 0, 25, 6), (23, 7, 25, 13),
    (26, 0, 27, 11), (26, 12, 27, 13),
    (28, 0, 40, 0), (28, 1, 40, 1), (28, 2, 40, 2), (28, 3, 40, 3),
    (28, 4, 40, 4), (28, 5, 40, 5), (28, 6, 40, 6), (28, 7, 40, 13),
    (41, 0, 45, 6), (41, 7, 45, 13),
    (46, 0, 46, 6), (46, 7, 46, 13), (47, 0, 48, 6), (47, 7, 48, 13),
]


def build_awb(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "AWB"
    ws.column_dimensions["A"].width = 11
    for c in range(1, 14):
        ws.column_dimensions[get_column_letter(c + 1)].width = 12
    for r, c, text in AWB_LABELS:
        label(ws, r, c, text)
    for box in AWB_BOXES:
        outline(ws, *box)
    for name, r1, c1, r2, c2 in AWB_FILLS:
        add_fill(ws, wb, "AWB", name, r1, c1, r2, c2)
    wb.save(path)


def main():
    os.makedirs(TEMPLATES_DIR, exist_ok=True)
    cmr = os.path.join(TEMPLATES_DIR, "cmr_template.xlsx")
    awb = os.path.join(TEMPLATES_DIR, "awb_template.xlsx")
    build_cmr(cmr)
    build_awb(awb)
    for p in (cmr, awb):
        print("wrote", os.path.normpath(p))


if __name__ == "__main__":
    main()
